# src/services/doc_generator.py
from typing import Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.services.deceased_service import get_address_info, get_case_by_id, get_deceased_by_case_id
from src.utils.date_utils import convert_seireki_to_wareki


# --- OXML ヘルパー: セルに特定の枠線を設定 ---
def set_cell_border(cell, **kwargs):
    """
    セルの枠線を設定する (タイトル枠用)
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for edge in ("left", "top", "right", "bottom", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)

            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


# --- ヘルパー: 数値を漢数字に変換 ---
def to_kanji_num(val) -> str:
    if val is None:
        return ""
    s = str(val)
    table = str.maketrans(
        {
            "0": "〇",
            "1": "一",
            "2": "二",
            "3": "三",
            "4": "四",
            "5": "五",
            "6": "六",
            "7": "七",
            "8": "八",
            "9": "九",
            ".": "・",
        }
    )
    return s.translate(table)


# --- フォント設定 ---
def set_font_style(run, font_name="MS Mincho", size=10.5, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font_name)


def generate_division_agreement_doc(
    case_id: int, allocations: Dict[str, int], output_path: str
) -> str:
    """
    遺産分割協議書 (20250413版 完全再現)
    """
    # 1. データ取得
    case = get_case_by_id(case_id)
    deceased = get_deceased_by_case_id(case_id)
    if not case or not deceased:
        raise ValueError("データ不足")

    heirs_map = {h.id: h for h in deceased.heirs}

    # 2. ドキュメント設定
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "MS Mincho"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "MS Mincho")
    style.paragraph_format.line_spacing = 1.15

    # ==========================================
    # 1. タイトル (枠囲み)
    # ==========================================
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_cell = title_table.cell(0, 0)

    p = title_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("遺産分割協議書")
    set_font_style(run, size=18, bold=True)

    border_style = {"val": "double", "sz": "4", "space": "0", "color": "000000"}
    set_cell_border(
        title_cell, top=border_style, bottom=border_style, left=border_style, right=border_style
    )

    title_table.autofit = False
    title_table.columns[0].width = Cm(12)

    doc.add_paragraph("")

    # ==========================================
    # 2. 被相続人情報
    # ==========================================
    d_addr = get_address_info("deceased", deceased.id)
    last_address = f"{d_addr.get('prefecture', '')}{d_addr.get('city_ward_town', '')}{d_addr.get('street_address', '')}"
    if d_addr.get("building_name"):
        last_address += f" {d_addr.get('building_name')}"
    last_hometown = deceased.hometown if deceased.hometown else "（未登録）"

    d_name = f"{deceased.name_last} {deceased.name_first}"
    death_date_str = (
        convert_seireki_to_wareki(deceased.date_of_death)
        if deceased.date_of_death
        else "平成/令和xx年xx月xx日"
    )

    def add_header_line(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        p.add_run(f"{label}　　").bold = False
        p.add_run(value)

    add_header_line("最後の住所　　", last_address)
    add_header_line("最後の本籍　　", last_hometown)
    add_header_line("登記簿上の住所", last_address)

    doc.add_paragraph("")

    # ==========================================
    # 3. 前文
    # ==========================================
    preamble = (
        f"被相続人　{d_name}　（{death_date_str}死亡）の遺産については、"
        "同人の相続人全員において分割協議を行った結果、各相続人がそれぞれ次の通り遺産を分割し、"
        "債務・葬式費用を負担することに決定した。"
    )
    p = doc.add_paragraph(preamble)
    p.paragraph_format.first_line_indent = Cm(0.5)

    note = (
        "なお、下記の不動産は住民票上の住所と不動産登記簿上の住所が一致していないが、"
        "全て被相続人の所有に係る不動産に間違いない。"
    )
    p = doc.add_paragraph(note)
    p.paragraph_format.first_line_indent = Cm(0.5)

    doc.add_paragraph("")
    p = doc.add_paragraph("記")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # ==========================================
    # 4. 財産目録
    # ==========================================
    assets_by_heir: Dict[int, Dict[str, List]] = {}

    for asset in case.real_estates:
        key = f"RE_{asset.id}"
        heir_id = allocations.get(key)
        if heir_id:
            if heir_id not in assets_by_heir:
                assets_by_heir[heir_id] = {"real_estates": [], "banks": [], "securities": []}
            assets_by_heir[heir_id]["real_estates"].append(asset)

    for asset in case.financial_assets:
        key = f"BANK_{asset.id}"
        heir_id = allocations.get(key)
        if heir_id:
            if heir_id not in assets_by_heir:
                assets_by_heir[heir_id] = {"real_estates": [], "banks": [], "securities": []}

            if asset.asset_type == "SECURITIES" or (
                asset.bank_ref and "証券" in asset.bank_ref.bank_name
            ):
                assets_by_heir[heir_id]["securities"].append(asset)
            else:
                assets_by_heir[heir_id]["banks"].append(asset)

    section_num = 1

    for heir_id, assets in assets_by_heir.items():
        heir = heirs_map.get(heir_id)
        if not heir:
            continue
        heir_name = f"{heir.name_last} {heir.name_first}"

        doc.add_paragraph(f"{section_num}．以下の財産は相続人　{heir_name}　が取得する。")

        sub_num = 1

        # --- (1) 不動産 ---
        if assets["real_estates"]:
            lands = [x for x in assets["real_estates"] if x.property_type == "Land"]
            buildings = [
                x for x in assets["real_estates"] if x.property_type in ["Building", "Condo"]
            ]

            if lands:
                doc.add_paragraph(f"（{sub_num}）土地")
                for re in lands:
                    _create_property_table(doc, re, "Land")
                    doc.add_paragraph(f"　被相続人の持分　　{re.ownership_share or '1/1'}")
                    doc.add_paragraph(
                        f"　相続人 {heir_name}が取得する持分　　被相続人所有のうち　1/1"
                    )
                    doc.add_paragraph("")
                sub_num += 1

            if buildings:
                doc.add_paragraph(f"（{sub_num}）建物")
                for re in buildings:
                    _create_property_table(doc, re, "Building")
                    doc.add_paragraph(f"　被相続人の持分　　{re.ownership_share or '1/1'}")
                    doc.add_paragraph(
                        f"　相続人 {heir_name}が取得する持分　　被相続人所有のうち　1/1"
                    )
                    doc.add_paragraph("")
                sub_num += 1

        # --- (2) 有価証券 ---
        if assets["securities"]:
            doc.add_paragraph(f"（{sub_num}）有価証券（未収配当金、未収分配金等の法定果実を含む）")

            headers = ["No", "種類", "銘柄", "所在場所等", "数量"]
            rows_data = []
            for idx, sec in enumerate(assets["securities"], 1):
                rows_data.append(
                    [
                        str(idx),
                        sec.account_type_ref.type_name if sec.account_type_ref else "",
                        "（銘柄不明）",
                        f"{sec.bank_ref.bank_name} {sec.branch_ref.branch_name if sec.branch_ref else ''}",
                        str(sec.balance),
                    ]
                )
            _create_list_table(doc, headers, rows_data)
            doc.add_paragraph("")
            sub_num += 1

        # --- (3) 現預金 ---
        if assets["banks"]:
            doc.add_paragraph(f"（{sub_num}）現預金")

            headers = ["No", "所在場所等", "支店", "種類", "口座番号"]
            col_widths = [Cm(1.0), Cm(6.0), Cm(4.0), Cm(2.5), Cm(4.0)]
            rows_data = []
            for idx, bank in enumerate(assets["banks"], 1):
                branch_str = bank.branch_ref.branch_name if bank.branch_ref else "－"
                if "ゆうちょ" in (bank.bank_ref.bank_name or ""):
                    branch_str = "－"

                rows_data.append(
                    [
                        str(idx),
                        bank.bank_ref.bank_name if bank.bank_ref else "",
                        branch_str,
                        bank.account_type_ref.type_name if bank.account_type_ref else "",
                        bank.account_number,
                    ]
                )
            _create_list_table(doc, headers, rows_data, widths=col_widths)
            doc.add_paragraph("")
            sub_num += 1

        section_num += 1

    # ==========================================
    # 5. 債務・葬式費用
    # ==========================================
    if case.liabilities:
        debts = [l for l in case.liabilities if l.is_debt and not l.is_funeral_cost]
        if debts:
            doc.add_paragraph(f"{section_num}．以下の債務は相続人　（代表者）　が負担する。")
            doc.add_paragraph("単位：円")
            headers = ["No", "種類", "細目", "債権者", "金額"]
            rows = []
            for idx, d in enumerate(debts, 1):
                rows.append([str(idx), "債務", d.description, "", f"{d.amount:,.0f}"])
            _create_list_table(doc, headers, rows)
            doc.add_paragraph("")
            section_num += 1

    doc.add_paragraph(f"{section_num}．葬式費用については相続人　（喪主）　が全額負担する。")
    doc.add_paragraph("")

    # ==========================================
    # 6. 後文・日付・署名
    # ==========================================
    doc.add_paragraph("")
    post_text = (
        "前記の通り相続人全員による遺産分割の協議が成立したので、これを証するため本書を作成し、"
        "以下に各自署名捺印する。なお、本協議書に記載なき遺産・債務並びに後日判明した遺産・債務は、"
        "相続人全員で別途協議して決めるものとする。"
    )
    doc.add_paragraph(post_text)
    doc.add_paragraph("")

    p = doc.add_paragraph("　　　　　年　　 月　　 日")
    doc.add_paragraph("")

    for heir in deceased.heirs:
        h_addr = get_address_info("heir", heir.id)
        address_str = f"{h_addr.get('prefecture', '')}{h_addr.get('city_ward_town', '')}{h_addr.get('street_address', '')}"
        if h_addr.get("building_name"):
            address_str += f" {h_addr.get('building_name')}"

        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(8)
        p.add_run(f"住所　　　{address_str}\n")
        p.add_run(f"相続人　　{heir.name_last} {heir.name_first}")

        doc.add_paragraph("")
        p = doc.add_paragraph("署名")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("")
        p = doc.add_paragraph("実印")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.right_indent = Cm(4)
        doc.add_paragraph("")
        doc.add_paragraph("")

    p = doc.add_paragraph("捨印")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.right_indent = Cm(2)

    try:
        doc.save(output_path)
    except Exception as e:
        if "Permission denied" in str(e):
            raise PermissionError("ファイルが開かれています。閉じてから再度実行してください。")
        raise e

    return output_path


def _create_property_table(doc, asset, p_type):
    """不動産の詳細を表示する表を作成 (枠線なし、項目名揃え)"""
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(11.0)

    def add_row(label, val):
        row = table.add_row()
        row.cells[0].text = label
        # 💡 修正点: Noneの場合は空文字に変換してエラーを防ぐ
        safe_val = str(val) if val is not None else ""
        row.cells[1].text = safe_val
        set_font_style(row.cells[0].paragraphs[0].runs[0])
        if row.cells[1].paragraphs[0].runs:
            set_font_style(row.cells[1].paragraphs[0].runs[0])

    if p_type == "Land":
        add_row("　　　　所在", asset.location)
        add_row("　　　地番", asset.lot_number)
        add_row("　　　　地目", asset.land_category or "宅地")
        area = to_kanji_num(asset.land_area)
        add_row("　　　　地積", f"{area} ㎡")
    else:
        add_row("　　　　所在", asset.location)
        add_row("　　　　家屋番号", asset.house_number)
        add_row("　　　　種類", asset.land_category or "居宅")
        add_row("　　　　構造", asset.structure)
        area = to_kanji_num(asset.floor_area)
        add_row("　　　　床面積", f"{area} ㎡")


def _create_list_table(doc, headers, data, widths=None):
    """一覧表を作成 (格子状の枠線あり)"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False

    if not widths:
        if len(headers) == 5:
            widths = [Cm(1.0), Cm(6.0), Cm(4.0), Cm(3.0), Cm(4.0)]
        else:
            widths = [Cm(18 / len(headers))] * len(headers)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i < len(widths):
            hdr_cells[i].width = widths[i]
        set_font_style(hdr_cells[i].paragraphs[0].runs[0], bold=True)

    for row_data in data:
        row = table.add_row()
        for i, cell_val in enumerate(row_data):
            # 💡 ここでも念のため str変換とNone対策
            safe_text = str(cell_val) if cell_val is not None else ""
            cell = row.cells[i]
            cell.text = safe_text

            if i < len(widths):
                cell.width = widths[i]

            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            if p.runs:
                set_font_style(p.runs[0], size=10.5)
